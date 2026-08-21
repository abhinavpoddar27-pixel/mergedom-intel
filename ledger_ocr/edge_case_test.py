"""
Edge-case test: mutate single input cells in a throw-away copy of the workbook,
recalculate with LibreOffice, and confirm the dependent formulas move.  This is
what distinguishes a live workbook from a static-looking transcription with
formula text pasted into it.
"""
import os
import shutil
import subprocess
import sys
import tempfile
from openpyxl import load_workbook

BOOK = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "Asian_Lakto_QC_Records_Aug2026.xlsx")
PROFILE = "file:///tmp/lo_profile"


def recalc(path, outdir):
    subprocess.run(["soffice", "--headless", "--norestore",
                    f"-env:UserInstallation={PROFILE}",
                    "--convert-to", "xlsx", "--outdir", outdir, path],
                   check=True, capture_output=True, timeout=400)
    return os.path.join(outdir, os.path.basename(path))


def probe(path):
    """Read the handful of roll-up numbers the checks depend on."""
    wb = load_workbook(path, data_only=True)
    idx = wb["Index & Checks"]
    out = {}
    for r in range(1, idx.max_row + 1):
        k = idx.cell(r, 1).value
        if isinstance(k, str) and k in (
                "Numeric torque readings transcribed",
                "Entries the chemist ringed as rejected",
                "Numeric readings the band formula flags as out of spec",
                "Band test disagrees with the chemist's ringing",
                "CSD :00/:30 reading pairs written on the forms",
                "CSD pairs whose two halves differ"):
            out[k] = idx.cell(r, 3).value
    ws = wb["Torque 20-08-2026"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "TOTAL" and isinstance(ws.cell(r, 2).value, (int, float)):
            out["T20 total readings"] = ws.cell(r, 2).value
            out["T20 minimum"] = round(ws.cell(r, 3).value, 2)
            out["T20 average"] = round(ws.cell(r, 5).value, 3)
    cs = wb["CSD 19-08-2026"]
    for r in range(1, cs.max_row + 1):
        if cs.cell(r, 1).value == "Pressure":
            out["CSD19 pressure min"] = cs.cell(r, 30).value    # AD
            out["CSD19 pressure max"] = cs.cell(r, 31).value    # AE
            out["CSD19 pressure avg"] = round(cs.cell(r, 32).value, 3)   # AF
        if cs.cell(r, 1).value == "Parameters filled in this column":
            out["CSD19 pairs differing"] = cs.cell(r, 34).value
    return out


def locate(ws, time_label, head):
    """Return the cell of the 'Reading' line for a given round and capper head."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == time_label and ws.cell(r, 2).value == "Reading (lb-in)":
            return ws.cell(r, 2 + head)
    raise LookupError(time_label)


CASES = [
    ("torque reading 9.8 -> 7.0 (head 5 was accepted; 7.0 is below the band)",
     lambda wb: setattr(locate(wb["Torque 20-08-2026"], "01:00", 5), "value", 7.0)),
    ("upper spec limit 19.0 -> 30.0 (should stop flagging most high readings)",
     lambda wb: setattr(wb["Index & Checks"]["B6"], "value", 30.0)),
    ("CSD 19-08 pressure at 09:30 42 -> 99 (breaks a :00/:30 pair)",
     lambda wb: setattr(wb["CSD 19-08-2026"]["F13"], "value", 99)),
]


def main():
    tmp = tempfile.mkdtemp(prefix="edgecase_")
    base_out = recalc(BOOK, os.path.join(tmp, "base"))
    base = probe(base_out)
    print("BASELINE")
    for k, v in base.items():
        print(f"   {k:<58} {v}")

    failures = 0
    for i, (label, mutate) in enumerate(CASES):
        print(f"\nCASE {i + 1}: {label}")
        work = os.path.join(tmp, f"c{i}.xlsx")
        shutil.copy(BOOK, work)
        wb = load_workbook(work)
        mutate(wb)
        wb.save(work)
        after = probe(recalc(work, os.path.join(tmp, f"out{i}")))
        moved = [k for k in base if base[k] != after.get(k)]
        for k in base:
            if base[k] != after.get(k):
                print(f"   CHANGED  {k:<56} {base[k]}  ->  {after[k]}")
        if not moved:
            print("   FAIL: nothing downstream moved - the workbook is not live")
            failures += 1
        else:
            print(f"   PASS: {len(moved)} dependent value(s) re-derived")

    shutil.rmtree(tmp, ignore_errors=True)
    print(f"\nedge-case tests: {len(CASES) - failures}/{len(CASES)} passed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
