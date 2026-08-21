"""
Post-build verification: load the LibreOffice-recalculated workbook and confirm
(a) no formula evaluated to an error, (b) every formula cell has a cached value,
(c) the verification roll-up reads as expected.
"""
import sys
from openpyxl import load_workbook

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "Err:", "#ERROR")

def main(recalced, original):
    src = load_workbook(original)                      # formulas as written
    out = load_workbook(recalced, data_only=True)      # values after recalc

    bad, empty, n_formula = [], [], 0
    for name in src.sheetnames:
        ws_f, ws_v = src[name], out[name]
        for row in ws_f.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formula += 1
                    v = ws_v[c.coordinate].value
                    if isinstance(v, str) and any(e in v for e in ERRORS):
                        bad.append(f"{name}!{c.coordinate} -> {v}   [{c.value[:70]}]")
                    elif v is None and '""' not in c.value:
                        # a formula containing "" may legitimately evaluate to blank
                        empty.append(f"{name}!{c.coordinate}   [{c.value[:70]}]")

    print(f"formula cells: {n_formula}")
    print(f"formula errors after recalculation: {len(bad)}")
    for b in bad:
        print("   ", b)
    print(f"formula cells with no cached value: {len(empty)}")
    for e in empty[:15]:
        print("   ", e)

    print("\n--- Index & Checks roll-up (recalculated) ---")
    idx = out["Index & Checks"]
    for r in range(1, idx.max_row + 1):
        a, b, c = idx.cell(r, 1).value, idx.cell(r, 3).value, idx.cell(r, 4).value
        if a and isinstance(a, str) and (
                "torque" in a.lower() or "CSD" in a or "Readings" in a or "Band test" in a):
            print(f"  {a:<52} {str(b):<8} (expected: {c})")

    print("\n--- Torque per-round statistics (recalculated) ---")
    for name in ("Torque 17-08-2026", "Torque 20-08-2026"):
        ws = out[name]
        print(f"  {name}")
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value in ("11:00", "01:00", "TOTAL") and isinstance(ws.cell(r, 2).value, (int, float)):
                vals = [ws.cell(r, c).value for c in range(1, 9)]
                print("     " + " | ".join(
                    f"{v:.2f}" if isinstance(v, float) else str(v) for v in vals))

    print("\n--- CSD per-sheet totals (recalculated) ---")
    for name in ("CSD 17-08-2026", "CSD 19-08-2026", "CSD 20-08-2026"):
        ws = out[name]
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == "Parameters filled in this column":
                print(f"  {name}: entries={ws.cell(r, 28).value}  "
                      f"pairs={ws.cell(r, 33).value}  pairs_differing={ws.cell(r, 34).value}")
    return 1 if (bad or empty) else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
