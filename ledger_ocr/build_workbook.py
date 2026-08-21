"""
Build the verified Excel workbook from the transcribed ledger data.

Every total, statistic and spec-check in the workbook is a live formula that
references the transcribed data cells, so correcting any single reading makes
the whole sheet re-derive itself.
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ledger_data as D

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Asian_Lakto_QC_Records_Aug2026.xlsx")

# ---------------------------------------------------------------- styling ---
FONT = "Calibri"
F_TITLE   = Font(name=FONT, size=15, bold=True, color="FFFFFF")
F_SUB     = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_H       = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_SECTION = Font(name=FONT, size=11, bold=True, color="1F3864")
F_LBL     = Font(name=FONT, size=10, bold=True)
F_B       = Font(name=FONT, size=10)
F_SMALL   = Font(name=FONT, size=9, italic=True, color="444444")
F_TOT     = Font(name=FONT, size=10, bold=True)
F_FLAG    = Font(name=FONT, size=10, bold=True, color="9C0006")

FILL_TITLE  = PatternFill("solid", fgColor="1F3864")
FILL_SUB    = PatternFill("solid", fgColor="2E5496")
FILL_HEAD   = PatternFill("solid", fgColor="4472C4")
FILL_BAND   = PatternFill("solid", fgColor="D9E2F3")
FILL_TOT    = PatternFill("solid", fgColor="FFF2CC")
FILL_META   = PatternFill("solid", fgColor="EDEDED")
FILL_BAD    = PatternFill("solid", fgColor="FFC7CE")
FILL_WARN   = PatternFill("solid", fgColor="FFE699")
FILL_GOOD   = PatternFill("solid", fgColor="E2EFDA")

_thin = Side(style="thin", color="9BA5B4")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

C_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
C_LEFTW  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
C_CTR    = Alignment(horizontal="center", vertical="center", wrap_text=True)
C_RIGHT  = Alignment(horizontal="right",  vertical="center")

FLAG_AUTHOR = "Transcription QA"


def page_setup(ws, title_rows=None):
    """Landscape, scaled to one page wide - these sheets are far wider than tall."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows


def put(ws, row, col, value, font=F_B, align=None, fill=None, border=True, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if align:
        c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def style_only(ws, row, col, font=None, align=None, fill=None, border=True):
    """Style a cell without writing to it (merged cells reject value writes)."""
    c = ws.cell(row=row, column=col)
    if font:
        c.font = font
    if align:
        c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    return c


def banner(ws, row, first_col, last_col, text, font=F_TITLE, fill=FILL_TITLE, height=22):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=first_col, value=text)
    c.font, c.fill, c.alignment = font, fill, C_CTR
    ws.row_dimensions[row].height = height
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = fill
    return c


def meta_pair(ws, row, col, label, value):
    put(ws, row, col, label, font=F_LBL, align=C_LEFT, fill=FILL_META)
    put(ws, row, col + 1, value, font=F_B, align=C_LEFT)


def flag(cell, text):
    cm = Comment(text, FLAG_AUTHOR)
    cm.width, cm.height = 340, 150
    cell.comment = cm
    cell.font = F_FLAG


# ============================================================== CSD sheets ===
T0 = 3                      # first time column (C)
NT = len(D.CSD_TIMES)       # 25
TL = T0 + NT - 1            # last time column (AA = 27)
C_ENTRIES, C_OKCNT, C_MIN, C_MAX, C_AVG, C_PAIRS, C_PDIFF = (TL + 1, TL + 2, TL + 3,
                                                             TL + 4, TL + 5, TL + 6, TL + 7)
LAST = C_PDIFF

# Display precision per numeric row, matching the convention the chemist uses on
# the form.  This pads trailing zeros only - no value is altered.
CSD_NUMFMT = {
    "As is brix (Bev.)": "0.00",
    "% Acidity": "0.00",
    "Pressure": "0",
    "Temperature": "0.0",
    "G.V.": "0.0",
}

CSD_FLAGS = {
    ("CSD 20-08-2026", "Temperature", 8): (
        "AMBIGUITY / SOURCE ANOMALY - 6.1 C.\n"
        "Digits are legible, but this breaks the :00/:30 pairing used everywhere else on "
        "the sheet (12:30 reads 10.1) and sits well below every other temperature logged "
        "that day (7.4-13.1). Transcribed exactly as written. Recommend verifying against "
        "the original paper record."),
    ("CSD 17-08-2026", "Temperature", 13): (
        "SOURCE ANOMALY - 12.8 C against 12.2 C at 02:00.\n"
        "The only :00/:30 pair on this sheet whose two halves differ. Both digits are "
        "clearly legible; reproduced as printed rather than harmonised. Worth a check "
        "against the original."),
    ("CSD 19-08-2026", "Temperature", 8): (
        "PARTIALLY OBSCURED - 14.2 C.\n"
        "The 'Filler stop due to outfeed knife setting & welding work' annotation is "
        "written across this cell and crosses the digits. The reading is supported by "
        "12:30, which reads 14.2 clearly."),
}


def build_csd(wb, rep):
    ws = wb.create_sheet(rep["sheet"])
    ws.sheet_properties.tabColor = "2E5496"

    banner(ws, 1, 1, LAST, rep["title"], F_TITLE, FILL_TITLE, 24)
    banner(ws, 2, 1, LAST, f"Doc No.: {rep['docno']}    |    transcribed from {rep['photo']}",
           F_SUB, FILL_SUB, 18)

    meta_pair(ws, 4, 1, "Date", rep["date"])
    meta_pair(ws, 5, 1, "Product", rep["product"])
    meta_pair(ws, 6, 1, "Chemist", rep["chemist"])
    meta_pair(ws, 7, 1, "Shift", rep["shift"])
    for i, lbl in enumerate(["Finished Syrup B.No.", "Control Drink As is Brix",
                             "Control Drink Invert Brix", "Control Drink % Acidity"]):
        rr = 4 + i
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=9)
        put(ws, rr, 5, lbl, F_LBL, C_LEFT, FILL_META)
        for k in range(5, 10):
            ws.cell(row=rr, column=k).fill = FILL_META
            ws.cell(row=rr, column=k).border = BORDER
        ws.merge_cells(start_row=rr, start_column=10, end_row=rr, end_column=13)
        put(ws, rr, 10, "(blank on form)", F_SMALL, C_LEFT)
        for k in range(10, 14):
            ws.cell(row=rr, column=k).border = BORDER

    hdr = 9
    put(ws, hdr, 1, "Parameter", F_H, C_CTR, FILL_HEAD)
    put(ws, hdr, 2, "Frequency", F_H, C_CTR, FILL_HEAD)
    for i, t in enumerate(D.CSD_TIMES):
        put(ws, hdr, T0 + i, t, F_H, C_CTR, FILL_HEAD)
        put(ws, hdr + 1, T0 + i, D.CSD_TIMES_24[i], F_SMALL, C_CTR, FILL_BAND)
    for col, name in ((C_ENTRIES, "Entries"), (C_OKCNT, "OK count"), (C_MIN, "Min"),
                      (C_MAX, "Max"), (C_AVG, "Average"), (C_PAIRS, ":00/:30 pairs"),
                      (C_PDIFF, "Pairs differing")):
        put(ws, hdr, col, name, F_H, C_CTR, FILL_HEAD)
        put(ws, hdr + 1, col, "", F_SMALL, C_CTR, FILL_BAND)
    put(ws, hdr + 1, 1, "24-hour equivalent →", F_SMALL, C_RIGHT, FILL_BAND)
    put(ws, hdr + 1, 2, "", F_SMALL, C_CTR, FILL_BAND)

    first = hdr + 2
    for r, (pname, freq, kind) in enumerate(D.CSD_PARAMS):
        row = first + r
        shade = FILL_BAND if r % 2 else None
        put(ws, row, 1, pname, F_LBL, C_LEFT, shade)
        put(ws, row, 2, freq, F_SMALL, C_LEFTW, shade)
        nf = CSD_NUMFMT.get(pname)
        for i in range(NT):
            v = rep["rows"][pname][i]
            cell = put(ws, row, T0 + i, v,
                       F_B, C_RIGHT if kind == "num" else C_CTR, shade,
                       fmt=nf if kind == "num" else None)
            key = (rep["sheet"], pname, i)
            if key in CSD_FLAGS:
                flag(cell, CSD_FLAGS[key])

        a, z = get_column_letter(T0), get_column_letter(TL)
        rng = f"${a}{row}:${z}{row}"
        pa, pz = get_column_letter(T0), get_column_letter(TL - 1)
        qa, qz = get_column_letter(T0 + 1), get_column_letter(TL)
        left, right = f"${pa}{row}:${pz}{row}", f"${qa}{row}:${qz}{row}"
        mask = f"(MOD(COLUMN({left})-COLUMN(${pa}$1),2)=0)"

        if kind == "num":
            put(ws, row, C_ENTRIES, f"=COUNT({rng})", F_TOT, C_RIGHT, shade)
            put(ws, row, C_OKCNT, "", F_B, C_CTR, shade)
            for col, fn in ((C_MIN, "MIN"), (C_MAX, "MAX"), (C_AVG, "AVERAGE")):
                put(ws, row, col, f'=IF(COUNT({rng})=0,"",{fn}({rng}))',
                    F_TOT, C_RIGHT, shade,
                    fmt="0.00" if fn == "AVERAGE" else (nf or "0.00"))
            put(ws, row, C_PAIRS,
                f'=SUMPRODUCT({mask}*({left}<>"")*({right}<>""))',
                F_TOT, C_RIGHT, shade)
            put(ws, row, C_PDIFF,
                f'=SUMPRODUCT({mask}*({left}<>"")*({right}<>"")*({left}<>{right}))',
                F_TOT, C_RIGHT, shade)
        else:
            put(ws, row, C_ENTRIES, f"=COUNTA({rng})", F_TOT, C_RIGHT, shade)
            put(ws, row, C_OKCNT, f'=COUNTIF({rng},"OK")', F_TOT, C_RIGHT, shade)
            for col in (C_MIN, C_MAX, C_AVG, C_PAIRS, C_PDIFF):
                put(ws, row, col, "", F_B, C_CTR, shade)

    last_p = first + len(D.CSD_PARAMS) - 1

    tot = last_p + 1
    put(ws, tot, 1, "Parameters filled in this column", F_TOT, C_LEFT, FILL_TOT)
    put(ws, tot, 2, "(live count)", F_SMALL, C_LEFT, FILL_TOT)
    for i in range(NT):
        cl = get_column_letter(T0 + i)
        put(ws, tot, T0 + i, f"=COUNTA({cl}{first}:{cl}{last_p})",
            F_TOT, C_CTR, FILL_TOT)
    a, z = get_column_letter(T0), get_column_letter(TL)
    put(ws, tot, C_ENTRIES, f"=SUM({a}{tot}:{z}{tot})", F_TOT, C_RIGHT, FILL_TOT)
    for col in (C_OKCNT, C_MIN, C_MAX, C_AVG):
        put(ws, tot, col, "", F_TOT, C_CTR, FILL_TOT)
    put(ws, tot, C_PAIRS, f"=SUM({get_column_letter(C_PAIRS)}{first}:"
                          f"{get_column_letter(C_PAIRS)}{last_p})", F_TOT, C_RIGHT, FILL_TOT)
    put(ws, tot, C_PDIFF, f"=SUM({get_column_letter(C_PDIFF)}{first}:"
                          f"{get_column_letter(C_PDIFF)}{last_p})", F_TOT, C_RIGHT, FILL_TOT)

    ws.conditional_formatting.add(
        f"{get_column_letter(C_PDIFF)}{first}:{get_column_letter(C_PDIFF)}{tot}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_WARN, font=F_FLAG))

    r = tot + 2
    banner(ws, r, 1, LAST, "Annotations written on the form", F_SUB, FILL_SUB, 18)
    r += 1
    put(ws, r, 1, "Where", F_H, C_CTR, FILL_HEAD)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST)
    put(ws, r, 2, "What the chemist wrote", F_H, C_LEFT, FILL_HEAD)
    for col in range(2, LAST + 1):
        ws.cell(row=r, column=col).fill = FILL_HEAD
    for where, what in rep["annotations"]:
        r += 1
        put(ws, r, 1, where, F_LBL, C_LEFTW)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST)
        put(ws, r, 2, what, F_B, C_LEFTW)
        ws.row_dimensions[r].height = 30
    if not rep["annotations"]:
        r += 1
        put(ws, r, 1, "(none)", F_SMALL, C_LEFT)

    r += 2
    put(ws, r, 1, "Remark (as written on form)", F_LBL, C_LEFT, FILL_META)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST)
    put(ws, r, 2, rep["remark"] or "(left blank on the form)", F_B, C_LEFTW)
    r += 1
    put(ws, r, 1, "QC Chemist", F_LBL, C_LEFT, FILL_META)
    put(ws, r, 2, rep["qc_chemist"], F_B, C_LEFT)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    put(ws, r, 5, "QC Manager", F_LBL, C_LEFT, FILL_META)
    for k in range(5, 9):
        ws.cell(row=r, column=k).fill = FILL_META
        ws.cell(row=r, column=k).border = BORDER
    ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=12)
    put(ws, r, 9, rep["qc_manager"], F_B, C_LEFT)
    for k in range(9, 13):
        ws.cell(row=r, column=k).border = BORDER
    if rep["note_below_table"]:
        r += 1
        put(ws, r, 1, "Note below the table", F_LBL, C_LEFTW, FILL_META)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LAST)
        put(ws, r, 2, rep["note_below_table"], F_B, C_LEFTW)
        ws.row_dimensions[r].height = 32

    r += 2
    banner(ws, r, 1, LAST, "Notes on this sheet", F_SUB, FILL_SUB, 18)
    notes = [
        "Blank cell = left blank on the paper form.   – = a dash was written "
        "(the test was not due at that interval).   OK = written as OK.",
        "The 24-hour row under the printed times is our addition; the form prints "
        "12-hour times with no am/pm marker.",
        "Numeric rows are displayed to a fixed number of decimals matching the form's own "
        "convention (brix and acidity 2 dp, temperature and G.V. 1 dp, pressure whole "
        "numbers). Where the chemist omitted a trailing zero the stored value is unchanged "
        "- only its display is padded.",
        "Entries / OK count / Min / Max / Average / pair columns are live formulas over "
        "the time columns to their left - correct any reading and they re-derive.",
        "“:00/:30 pairs” counts the paired readings actually written; "
        "“Pairs differing” counts pairs whose two halves are not identical. "
        "That is the accuracy check for this form, which carries no arithmetic column.",
        "Cells carrying a red figure have a comment attached explaining the ambiguity "
        "or source anomaly.",
    ]
    own = [(pname, i, txt) for (sh, pname, i), txt in CSD_FLAGS.items()
           if sh == rep["sheet"]]
    if own:
        for pname, i, txt in sorted(own, key=lambda x: x[1]):
            notes.append(f"FLAGGED - {pname} at {D.CSD_TIMES[i]}: "
                         + " ".join(txt.split("\n")[1:]).strip())
    else:
        notes.append("No cell on this sheet required an ambiguity flag.")
    for n in notes:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LAST)
        put(ws, r, 1, "•  " + n, F_SMALL, C_LEFTW)
        ws.row_dimensions[r].height = 26

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    for i in range(NT):
        ws.column_dimensions[get_column_letter(T0 + i)].width = 7.6
    for col in (C_ENTRIES, C_OKCNT, C_MIN, C_MAX, C_AVG, C_PAIRS, C_PDIFF):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.freeze_panes = ws.cell(row=first, column=T0)
    ws.sheet_view.showGridLines = False
    page_setup(ws, title_rows=f"{hdr}:{hdr + 1}")
    return ws


# =========================================================== Torque sheets ===
TH0 = 3                       # first capper-head column (C)
NH = 15
THL = TH0 + NH - 1            # Q
C_REMARK = THL + 1            # R
C_STAT   = THL + 2            # S
TLAST    = C_STAT

TORQUE_FLAGS = {
    ("Torque 20-08-2026", "11:00", 6): (
        "SOURCE GAP - 7.9 lb-in is ringed as rejected, but unlike every other ringed "
        "reading on either sheet no re-test value was written underneath. The form "
        "leaves this head's outcome unrecorded."),
    ("Torque 20-08-2026", "01:00", 11): (
        "AMBIGUOUS DIGITS - read as 20.0. The last character is heavily over-written and "
        "could also be 20.6. It is ringed as rejected either way, and the re-test (11.5) "
        "is recorded, so the outcome is unaffected. Verify against the original."),
    ("Torque 20-08-2026", "01:00", 6): (
        "AMBIGUOUS DIGIT - read as 13.2. The final digit is written cursively and could "
        "be 6 (13.6). Both readings are inside the acceptance band, so the 'OK' result "
        "is unaffected."),
    ("Torque 20-08-2026", "11:00", 7): (
        "NON-NUMERIC ENTRY - the chemist wrote 'No open' (the cap would not release), "
        "with 'open' written above the cell. A re-test of 16.6 lb-in was recorded."),
    ("Torque 17-08-2026", "11:00", 14): (
        "NON-NUMERIC ENTRY - the chemist wrote 'No open', with 'open' written above the "
        "cell. A re-test of 16.0 lb-in was recorded."),
    ("Torque 17-08-2026", "11:00", 10): (
        "AMBIGUOUS DIGIT - read as 6.0. The leading digit is written over itself; 6 is "
        "the better-supported reading (a 0 is also conceivable). It is ringed as rejected "
        "either way and re-tested at 10.8."),
    ("Torque 17-08-2026", "01:00", 13): (
        "PARTIALLY OVERLAPPED - read as 25.1. The final digit sits on the cell boundary "
        "and is crossed by the ring drawn around the neighbouring head-14 reading."),
}


def build_torque(wb, rep):
    ws = wb.create_sheet(rep["sheet"])
    ws.sheet_properties.tabColor = "C55A11"

    banner(ws, 1, 1, TLAST, rep["company"], F_TITLE, FILL_TITLE, 24)
    banner(ws, 2, 1, TLAST, rep["title"], F_SUB, FILL_SUB, 20)
    banner(ws, 3, 1, TLAST, f"transcribed from {rep['photo']}", F_SMALL,
           PatternFill("solid", fgColor="EDEDED"), 15)

    meta_pair(ws, 5, 1, "DATE", rep["date"])
    meta_pair(ws, 6, 1, "PRODUCT", rep["product"])
    meta_pair(ws, 7, 1, "CAP MAKE", rep["capmake"])
    meta_pair(ws, 5, 5, "SHIFT", rep["shift"])
    meta_pair(ws, 6, 5, "CHEMIST NAME", rep["chemist"])

    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=TLAST)
    put(ws, 8, 1,
        "Acceptance band used by the spec checks below: "
        "=\"\"&TEXT(SPEC_LOW,\"0.0\")", F_SMALL, C_LEFTW)
    ws.cell(row=8, column=1).value = ('=\"Acceptance band used by the checks below: \"&'
                                      'TEXT(SPEC_LOW,"0.0")&" to "&TEXT(SPEC_HIGH,"0.0")&'
                                      '" lb-in  (editable on the Index & Checks sheet - '
                                      'the form itself prints no limits)"')

    hdr = 10
    put(ws, hdr, 1, "TIME", F_H, C_CTR, FILL_HEAD)
    put(ws, hdr, 2, "Entry", F_H, C_CTR, FILL_HEAD)
    for h in range(NH):
        put(ws, hdr, TH0 + h, h + 1, F_H, C_CTR, FILL_HEAD)
    put(ws, hdr, C_REMARK, "REMARK", F_H, C_CTR, FILL_HEAD)
    put(ws, hdr, C_STAT, "", F_H, C_CTR, FILL_HEAD)
    ws.merge_cells(start_row=hdr - 1, start_column=TH0, end_row=hdr - 1, end_column=THL)
    put(ws, hdr - 1, TH0, "CAPPER HEAD", F_H, C_CTR, FILL_SUB)
    for col in range(TH0, THL + 1):
        ws.cell(row=hdr - 1, column=col).fill = FILL_SUB
    put(ws, hdr - 1, 1, "", F_H, C_CTR, FILL_SUB)
    put(ws, hdr - 1, 2, "", F_H, C_CTR, FILL_SUB)
    put(ws, hdr - 1, C_REMARK, "", F_H, C_CTR, FILL_SUB)
    put(ws, hdr - 1, C_STAT, "", F_H, C_CTR, FILL_SUB)

    row = hdr + 1
    read_row, ring_row = {}, {}
    for t in D.TORQUE_TIMES:
        heads = rep["readings"].get(t)
        shade = None if heads else FILL_META
        put(ws, row, 1, t, F_LBL, C_CTR, shade)
        put(ws, row, 2, "Reading (lb-in)", F_SMALL, C_LEFT, shade)
        put(ws, row + 1, 1, "", F_LBL, C_CTR, shade)
        put(ws, row + 1, 2, "Re-test / Result", F_SMALL, C_LEFT, shade)
        for h in range(NH):
            v = heads[h][0] if heads else None
            rt = heads[h][2] if heads else None
            c1 = put(ws, row, TH0 + h, v, F_B,
                     C_RIGHT if isinstance(v, (int, float)) else C_CTR, shade,
                     fmt="0.0" if isinstance(v, (int, float)) else None)
            put(ws, row + 1, TH0 + h, rt, F_B,
                C_RIGHT if isinstance(rt, (int, float)) else C_CTR, shade,
                fmt="0.0" if isinstance(rt, (int, float)) else None)
            key = (rep["sheet"], t, h + 1)
            if key in TORQUE_FLAGS:
                flag(c1, TORQUE_FLAGS[key])
        put(ws, row, C_REMARK, rep["remarks"].get(t, ""), F_B, C_LEFTW, shade)
        put(ws, row + 1, C_REMARK, "", F_B, C_LEFTW, shade)
        rng = f"${get_column_letter(TH0)}{row}:${get_column_letter(THL)}{row}"
        if heads:
            put(ws, row, C_STAT, f"=COUNT({rng})&\" numeric\"", F_SMALL, C_LEFT, shade)
            read_row[t] = row
        else:
            put(ws, row, C_STAT, "not recorded", F_SMALL, C_LEFT, shade)
        put(ws, row + 1, C_STAT, "", F_SMALL, C_LEFT, shade)
        row += 2
    end_tbl = row - 1

    # ---------------------------------------------------------- verification
    row += 1
    banner(ws, row, 1, TLAST, "Verification - every cell below is a live formula",
           F_SUB, FILL_SUB, 18)
    vhdr = row + 1
    put(ws, vhdr, 1, "TIME", F_H, C_CTR, FILL_HEAD)
    put(ws, vhdr, 2, "Check", F_H, C_CTR, FILL_HEAD)
    for h in range(NH):
        put(ws, vhdr, TH0 + h, h + 1, F_H, C_CTR, FILL_HEAD)
    put(ws, vhdr, C_REMARK, "Count", F_H, C_CTR, FILL_HEAD)
    put(ws, vhdr, C_STAT, "", F_H, C_CTR, FILL_HEAD)

    row = vhdr + 1
    mismatch_cells, out_cells, ring_cells = [], [], []
    for t in D.TORQUE_TIMES:
        heads = rep["readings"].get(t)
        if not heads:
            continue
        src = row
        put(ws, src, 1, t, F_LBL, C_CTR, FILL_BAND)
        put(ws, src, 2, "Ringed on the form (as transcribed)", F_SMALL, C_LEFT, FILL_BAND)
        put(ws, src + 1, 1, "", F_LBL, C_CTR)
        put(ws, src + 1, 2, "Outside band (formula)", F_SMALL, C_LEFT)
        put(ws, src + 2, 1, "", F_LBL, C_CTR, FILL_BAND)
        put(ws, src + 2, 2, "Formula agrees with the form?", F_SMALL, C_LEFT, FILL_BAND)
        rr = read_row[t]
        for h in range(NH):
            cl = get_column_letter(TH0 + h)
            put(ws, src, TH0 + h, "Y" if heads[h][1] else "", F_B, C_CTR, FILL_BAND)
            put(ws, src + 1, TH0 + h,
                f'=IF(NOT(ISNUMBER({cl}{rr})),"n/a",'
                f'IF(OR({cl}{rr}<SPEC_LOW,{cl}{rr}>SPEC_HIGH),"OUT","in"))',
                F_B, C_CTR)
            put(ws, src + 2, TH0 + h,
                f'=IF(NOT(ISNUMBER({cl}{rr})),"n/a",'
                f'IF(({cl}{src}="Y")=OR({cl}{rr}<SPEC_LOW,{cl}{rr}>SPEC_HIGH),'
                f'"agrees","MISMATCH"))',
                F_B, C_CTR, FILL_BAND)
        a, z = get_column_letter(TH0), get_column_letter(THL)
        put(ws, src, C_REMARK, f'=COUNTIF(${a}{src}:${z}{src},"Y")', F_TOT, C_CTR, FILL_BAND)
        put(ws, src + 1, C_REMARK, f'=COUNTIF(${a}{src+1}:${z}{src+1},"OUT")', F_TOT, C_CTR)
        put(ws, src + 2, C_REMARK,
            f'=COUNTIF(${a}{src+2}:${z}{src+2},"MISMATCH")', F_TOT, C_CTR, FILL_BAND)
        put(ws, src, C_STAT, "rejected by the chemist", F_SMALL, C_LEFT, FILL_BAND)
        put(ws, src + 1, C_STAT, "flagged by the formula", F_SMALL, C_LEFT)
        put(ws, src + 2, C_STAT, "must be 0", F_SMALL, C_LEFT, FILL_BAND)
        ring_cells.append(f"{get_column_letter(C_REMARK)}{src}")
        out_cells.append(f"{get_column_letter(C_REMARK)}{src+1}")
        mismatch_cells.append(f"{get_column_letter(C_REMARK)}{src+2}")
        ws.conditional_formatting.add(
            f"${a}{src+2}:${z}{src+2}",
            CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=FILL_BAD, font=F_FLAG))
        ws.conditional_formatting.add(
            f"${a}{src+1}:${z}{src+1}",
            CellIsRule(operator="equal", formula=['"OUT"'], fill=FILL_WARN))
        row = src + 3

    # ------------------------------------------------------------ statistics
    row += 1
    banner(ws, row, 1, 8, "Statistics per test round - live formulas", F_SUB, FILL_SUB, 18)
    shdr = row + 1
    stat_cols = ["Test round", "Numeric readings", "Minimum", "Maximum", "Average",
                 "Rejected (ringed)", "Re-tests written", "Re-tests inside band"]
    for i, h in enumerate(stat_cols):
        put(ws, shdr, 1 + i, h, F_H, C_CTR, FILL_HEAD)
    row = shdr + 1
    stat_rows = []
    for t in D.TORQUE_TIMES:
        if t not in read_row:
            continue
        rr = read_row[t]
        a, z = get_column_letter(TH0), get_column_letter(THL)
        rng = f"${a}{rr}:${z}{rr}"
        rtr = f"${a}{rr+1}:${z}{rr+1}"
        vsrc = [c for c in ring_cells if True]
        put(ws, row, 1, t, F_LBL, C_CTR, FILL_TOT)
        put(ws, row, 2, f"=COUNT({rng})", F_TOT, C_CTR, FILL_TOT)
        put(ws, row, 3, f"=MIN({rng})", F_TOT, C_CTR, FILL_TOT, fmt="0.0")
        put(ws, row, 4, f"=MAX({rng})", F_TOT, C_CTR, FILL_TOT, fmt="0.0")
        put(ws, row, 5, f"=AVERAGE({rng})", F_TOT, C_CTR, FILL_TOT, fmt="0.00")
        srow = [c for c in mismatch_cells]
        put(ws, row, 6,
            f'=COUNTIF(${a}{vhdr + 1 + 3 * list(read_row).index(t)}:'
            f'${z}{vhdr + 1 + 3 * list(read_row).index(t)},"Y")', F_TOT, C_CTR, FILL_TOT)
        put(ws, row, 7, f"=COUNT({rtr})", F_TOT, C_CTR, FILL_TOT)
        put(ws, row, 8,
            f'=SUMPRODUCT(({rtr}<>"")*(ISNUMBER({rtr}))*({rtr}>=SPEC_LOW)*({rtr}<=SPEC_HIGH))',
            F_TOT, C_CTR, FILL_TOT)
        stat_rows.append(row)
        row += 1

    put(ws, row, 1, "TOTAL", F_TOT, C_CTR, FILL_TOT)
    a, z = get_column_letter(TH0), get_column_letter(THL)
    all_reads = ",".join(f"${a}{read_row[t]}:${z}{read_row[t]}" for t in read_row)
    for col, fn in ((2, "SUM"), (3, "MIN"), (4, "MAX"), (5, "AVERAGE"),
                    (6, "SUM"), (7, "SUM"), (8, "SUM")):
        cl = get_column_letter(col)
        refs = ",".join(f"{cl}{r}" for r in stat_rows)
        # the overall mean must come from the readings themselves, not from the
        # per-round means, which would weight unequal round sizes equally.
        expr = f"=AVERAGE({all_reads})" if col == 5 else f"={fn}({refs})"
        put(ws, row, col, expr, F_TOT, C_CTR, FILL_TOT,
            fmt="0.00" if col == 5 else ("0.0" if col in (3, 4) else None))
    total_stat_row = row

    row += 2
    banner(ws, row, 1, TLAST, "Notes on this sheet", F_SUB, FILL_SUB, 18)
    notes = [
        "The paper form prints six test rounds (09:00, 11:00, 01:00, 03:00, 05:00, 07:00). "
        "Only 11:00 and 01:00 were filled in; the greyed rounds were left blank and are "
        "reproduced empty rather than dropped.",
        "A reading the chemist ringed on the form was rejected, and the re-test he then "
        "wrote underneath appears on the 'Re-test / Result' line. 'OK' on that line means "
        "the first reading was accepted.",
        "The form prints no acceptance limits. The band on the Index sheet is the one "
        "consistent with every ringing decision made across both sheets - see that sheet "
        "for how it was derived. Change the two cells and every check here re-derives.",
        "'Formula agrees with the form?' compares the band test against the chemist's own "
        "ringing for each head. Any MISMATCH means a digit was probably misread - the "
        "count must be 0.",
        "Cells carrying a red figure have a comment attached explaining the ambiguity.",
    ]
    own = [(t, h, txt) for (sh, t, h), txt in TORQUE_FLAGS.items() if sh == rep["sheet"]]
    if own:
        for t, h, txt in sorted(own):
            notes.append(f"FLAGGED - {t}, capper head {h}: " + " ".join(txt.split()))
    else:
        notes.append("No cell on this sheet required an ambiguity flag.")
    for n in notes:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=TLAST)
        put(ws, row, 1, "•  " + n, F_SMALL, C_LEFTW)
        ws.row_dimensions[row].height = 26

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 29
    for h in range(NH):
        ws.column_dimensions[get_column_letter(TH0 + h)].width = 8.4
    ws.column_dimensions[get_column_letter(C_REMARK)].width = 14
    ws.column_dimensions[get_column_letter(C_STAT)].width = 24
    ws.freeze_panes = ws.cell(row=hdr + 1, column=TH0)
    ws.sheet_view.showGridLines = False
    page_setup(ws, title_rows=f"{hdr}:{hdr}")
    return ws, mismatch_cells, out_cells, ring_cells, total_stat_row


# ================================================================== Index ====
def build_index(wb, csd_sheets, torque_meta):
    ws = wb.create_sheet("Index & Checks", 0)
    ws.sheet_properties.tabColor = "375623"
    W = 9

    banner(ws, 1, 1, W, "ASIAN LAKTO INDUSTRIES LTD – QC RECORDS, AUGUST 2026",
           F_TITLE, FILL_TITLE, 26)
    banner(ws, 2, 1, W,
           "Five photographed paper forms, transcribed and verified. "
           "One sheet per source document.", F_SUB, FILL_SUB, 18)

    r = 4
    put(ws, r, 1, "Acceptance band for cap-removal torque (editable)", F_SECTION, C_LEFT, border=False)
    r += 1
    put(ws, r, 1, "Lower limit (lb-in)", F_LBL, C_LEFT, FILL_META)
    put(ws, r, 2, D.SPEC_LOW_DEFAULT, F_TOT, C_CTR, FILL_TOT, fmt="0.0")
    low_ref = f"'Index & Checks'!$B${r}"
    r += 1
    put(ws, r, 1, "Upper limit (lb-in)", F_LBL, C_LEFT, FILL_META)
    put(ws, r, 2, D.SPEC_HIGH_DEFAULT, F_TOT, C_CTR, FILL_TOT, fmt="0.0")
    high_ref = f"'Index & Checks'!$B${r}"
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=W)
    put(ws, r, 1,
        "The torque form prints no specification limits. These two figures were derived "
        "from the source itself: across both torque sheets the chemist ringed 19 readings "
        "as rejected and left 39 un-ringed. Every un-ringed reading falls in 9.6–18.5 "
        "lb-in and every ringed one falls outside 7.9–20.0, so the true limits must lie "
        "between those bounds. 8.0 and 19.0 sit inside that gap. Replace them with the "
        "plant's actual specification and every check in the workbook re-derives.",
        F_SMALL, C_LEFTW)
    for k in range(3):
        ws.row_dimensions[r + k].height = 15
    r += 4

    put(ws, r, 1, "Sheets in this workbook", F_SECTION, C_LEFT, border=False)
    r += 1
    heads = ["Sheet", "Source photo", "Report", "Date", "Chemist", "Shift",
             "Product", "Data captured"]
    for i, h in enumerate(heads):
        put(ws, r, 1 + i, h, F_H, C_CTR, FILL_HEAD)
    r += 1
    for rep in D.CSD_REPORTS:
        put(ws, r, 1, rep["sheet"], F_LBL, C_LEFT)
        put(ws, r, 2, rep["photo"], F_B, C_LEFT)
        put(ws, r, 3, rep["title"], F_B, C_LEFT)
        put(ws, r, 4, rep["date"], F_B, C_CTR)
        put(ws, r, 5, rep["chemist"], F_B, C_LEFT)
        put(ws, r, 6, rep["shift"], F_B, C_CTR)
        put(ws, r, 7, rep["product"], F_B, C_LEFT)
        put(ws, r, 8, f"='{rep['sheet']}'!{get_column_letter(C_ENTRIES)}"
                      f"{11 + len(D.CSD_PARAMS)}&\" cell entries\"", F_B, C_LEFT)
        r += 1
    for rep, meta in zip(D.TORQUE_REPORTS, torque_meta):
        put(ws, r, 1, rep["sheet"], F_LBL, C_LEFT)
        put(ws, r, 2, rep["photo"], F_B, C_LEFT)
        put(ws, r, 3, rep["title"], F_B, C_LEFT)
        put(ws, r, 4, rep["date"], F_B, C_CTR)
        put(ws, r, 5, rep["chemist"], F_B, C_LEFT)
        put(ws, r, 6, rep["shift"], F_B, C_CTR)
        put(ws, r, 7, f"{rep['product']}  (cap make {rep['capmake']})", F_B, C_LEFT)
        put(ws, r, 8, f"='{rep['sheet']}'!$B${meta[4]}&\" torque readings\"", F_B, C_LEFT)
        r += 1

    r += 1
    put(ws, r, 1, "Live verification roll-up", F_SECTION, C_LEFT, border=False)
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    put(ws, r, 1, "Check", F_H, C_CTR, FILL_HEAD)
    style_only(ws, r, 2, F_H, C_CTR, FILL_HEAD)
    put(ws, r, 3, "Result", F_H, C_CTR, FILL_HEAD)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    put(ws, r, 4, "Expected", F_H, C_CTR, FILL_HEAD)
    for k in (5, 6):
        style_only(ws, r, k, F_H, C_CTR, FILL_HEAD)
    r += 1

    def rollup(label, formula, expected, good=True):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        put(ws, r, 1, label, F_LBL, C_LEFT)
        style_only(ws, r, 2)
        c = put(ws, r, 3, formula, F_TOT, C_CTR, FILL_GOOD if good else FILL_TOT)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        put(ws, r, 4, expected, F_SMALL, C_LEFTW)
        for k in (5, 6):
            style_only(ws, r, k)
        r += 1
        return c

    tq = [m for m in torque_meta]
    mism = "+".join(f"'{rep['sheet']}'!{c}" for rep, m in zip(D.TORQUE_REPORTS, tq) for c in m[0])
    outs = "+".join(f"'{rep['sheet']}'!{c}" for rep, m in zip(D.TORQUE_REPORTS, tq) for c in m[1])
    rings = "+".join(f"'{rep['sheet']}'!{c}" for rep, m in zip(D.TORQUE_REPORTS, tq) for c in m[2])
    reads = "+".join(f"'{rep['sheet']}'!$B${m[4]}" for rep, m in zip(D.TORQUE_REPORTS, tq))

    rollup("Numeric torque readings transcribed", f"={reads}", "58")
    rollup("Entries the chemist ringed as rejected", f"={rings}",
           "21 = 19 numeric readings + 2 'No open' entries")
    rollup("Numeric readings the band formula flags as out of spec", f"={outs}",
           "19 – must equal the ringed numeric readings")
    rollup("Band test disagrees with the chemist's ringing", f"={mism}",
           "0 – any other value means a digit was misread")

    csd_pairs = "+".join(f"'{s}'!{get_column_letter(C_PAIRS)}{11 + len(D.CSD_PARAMS)}"
                         for s in csd_sheets)
    csd_diff = "+".join(f"'{s}'!{get_column_letter(C_PDIFF)}{11 + len(D.CSD_PARAMS)}"
                        for s in csd_sheets)
    rollup("CSD :00/:30 reading pairs written on the forms", f"={csd_pairs}", "102")
    rollup("CSD pairs whose two halves differ", f"={csd_diff}",
           "2 – both are source anomalies, flagged with comments")

    r += 1
    put(ws, r, 1, "Flags, ambiguities and source anomalies", F_SECTION, C_LEFT, border=False)
    r += 1
    for i, h in enumerate(["Sheet", "Where", "Kind", "What we found and what we did",
                           "", "", "", "", ""]):
        put(ws, r, 1 + i, h, F_H, C_CTR, FILL_HEAD)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=W)
    r += 1

    FLAGS = [
        ("Torque 20-08-2026", "11:00, head 6", "Source gap",
         "7.9 lb-in is ringed as rejected but no re-test was written underneath, unlike "
         "every other ringed reading on either sheet. The head's outcome is unrecorded on "
         "the form; the cell is left empty here rather than filled in."),
        ("CSD 20-08-2026", "Temperature, 12:00", "Source anomaly / verify",
         "Reads 6.1 °C. Legible, but it breaks the :00/:30 pairing used throughout the "
         "sheet (12:30 reads 10.1) and is far below every other temperature logged that day. "
         "Reproduced as written."),
        ("CSD 17-08-2026", "Temperature, 02:30", "Source anomaly / verify",
         "Reads 12.8 °C against 12.2 °C at 02:00 – the only unequal pair on the "
         "sheet. Both digits are clear; reproduced rather than harmonised."),
        ("Torque 20-08-2026", "01:00, head 11", "Ambiguous digits",
         "Read as 20.0; the last character is over-written and could be 20.6. Ringed as "
         "rejected either way and re-tested at 11.5, so the outcome is unaffected."),
        ("Torque 20-08-2026", "01:00, head 6", "Ambiguous digit",
         "Read as 13.2; the final digit could be 6. Both values are inside the band, so "
         "the OK result is unaffected."),
        ("Torque 17-08-2026", "11:00, head 10", "Ambiguous digit",
         "Read as 6.0; the leading digit is written over itself and 0 is conceivable. "
         "Ringed and re-tested at 10.8 either way."),
        ("Torque 17-08-2026", "01:00, head 13", "Partially overlapped",
         "Read as 25.1; the last digit sits on the cell boundary and is crossed by the ring "
         "drawn around head 14."),
        ("Torque 17-08-2026 / 20-08-2026", "11:00, heads 14 and 7", "Non-numeric entry",
         "'No open' written instead of a reading (the cap would not release), with 'open' "
         "written above the cell. Re-tests of 16.0 and 16.6 lb-in were recorded. Kept as "
         "text; the statistics count numeric readings only."),
        ("CSD 17-08-2026 and 20-08-2026", "'As is brix' and '% Acidity' rows", "Source gap",
         "Both rows are entirely blank although the form requires brix every 30 minutes and "
         "acidity every hour. The 19-08 sheet has them filled. Left blank here."),
        ("CSD 20-08-2026", "Note below the table", "Cut off in the photograph",
         "The handwritten note ends '...time 7:00 Pm start' and the final line runs off the "
         "bottom edge of the photograph. Transcribed as far as it is visible."),
        ("CSD 17-08-2026", "Filler start annotation", "Uncertain word",
         "Read as 'Filler Start 8:38 AM'. The verb is written indistinctly; 'Start' is "
         "supported by readings beginning at 09:00 and by the same phrase on the 19-08 sheet."),
        ("All CSD sheets", "Header block and lot-number rows", "Blank on form",
         "Finished Syrup B.No., Control Drink As is Brix / Invert Brix / % Acidity, and the "
         "Closure, Label and Preform Lot No. rows were all left blank by the chemist."),
        ("All sheets", "Product name", "Reading note",
         "Read as 'MC2 Original 250 ml'; the 2 is written small, close to a subscript."),
        ("Both torque sheets", "Acceptance limits", "Not printed on the form",
         "The band used by the checks was inferred from the chemist's own ringing – see "
         "the top of this sheet. It is not a figure taken from the document."),
    ]
    for sheet, where, kind, what in FLAGS:
        put(ws, r, 1, sheet, F_B, C_LEFTW)
        put(ws, r, 2, where, F_B, C_LEFTW)
        put(ws, r, 3, kind, F_LBL, C_LEFTW)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=W)
        put(ws, r, 4, what, F_B, C_LEFTW)
        for k in range(4, W + 1):
            ws.cell(row=r, column=k).border = BORDER
        ws.row_dimensions[r].height = 42
        r += 1

    r += 1
    put(ws, r, 1, "How this workbook was produced", F_SECTION, C_LEFT, border=False)
    r += 1
    method = [
        "Five photographs were extracted from the supplied PDF at their embedded resolution "
        "(no re-rendering) and checked for EXIF rotation; none needed rotating.",
        "The photographs group into two report families: three dated CSD PET Online Record "
        "sheets (17, 19 and 20 August) and two dated Torque Testing Report sheets (17 and "
        "20 August). No photograph is a continuation page of another, so each gets its own sheet.",
        "Column and row rulings were detected computationally in each photograph, then every "
        "data cell was cut out individually and read at high magnification, so no value "
        "depends on judging which column a digit belongs to by eye.",
        "Neither form carries an arithmetic column to check against, so two structural "
        "checks were used instead: on the torque sheets, whether a single acceptance band "
        "separates the readings the chemist ringed from the ones he did not (58 of 58 agree); "
        "on the CSD sheets, whether the :00 and :30 halves of each logged pair match "
        "(100 of 102 do, and both exceptions are listed above).",
        "Every total, minimum, maximum, average, count and spec test in this workbook is a "
        "formula over the transcribed cells. The workbook was recalculated with LibreOffice "
        "and contains no formula errors.",
        "Where the paper form groups rows with merged cells or repeated headers, this "
        "workbook uses a single header row plus a shaded banner instead. Times are printed "
        "on the form in 12-hour form with no am/pm marker; a 24-hour row was added beneath "
        "them on the CSD sheets. Both are presentation changes only – no value was altered.",
    ]
    for m in method:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=W)
        put(ws, r, 1, "•  " + m, F_SMALL, C_LEFTW)
        ws.row_dimensions[r].height = 40
        r += 1

    widths = [30, 23, 31, 12, 17, 8, 35, 22, 16]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[get_column_letter(1 + i)].width = wdt
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"
    page_setup(ws)
    return low_ref, high_ref


def main():
    wb = Workbook()
    wb.remove(wb.active)

    csd_sheets = []
    for rep in D.CSD_REPORTS:
        build_csd(wb, rep)
        csd_sheets.append(rep["sheet"])

    torque_meta = []
    for rep in D.TORQUE_REPORTS:
        _ws, mism, outs, rings, tot = build_torque(wb, rep)
        torque_meta.append((mism, outs, rings, None, tot))

    low_ref, high_ref = build_index(wb, csd_sheets, torque_meta)
    wb.defined_names.add(DefinedName("SPEC_LOW", attr_text=low_ref))
    wb.defined_names.add(DefinedName("SPEC_HIGH", attr_text=high_ref))

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
